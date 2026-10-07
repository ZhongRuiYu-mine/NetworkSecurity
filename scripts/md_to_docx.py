# -*- coding: utf-8 -*-
"""把 results/method_analysis.md 转成 docx（报告章节可直接贴进 Word）。

只支持本仓库报告里实际用到的 Markdown 子集：
    # / ## / ### / #### 标题、| 表格 |、* 与 - 列表、> 引用、``` 代码块、
    **粗体**、`code`、--- 分隔线。
不追求通用性——通用转换用 pandoc，本脚本是为了在没有 pandoc 的机器上也能出 docx。

用法：
    python scripts/md_to_docx.py                       # 默认转 method_analysis.md
    python scripts/md_to_docx.py in.md out.docx
"""
from __future__ import annotations

import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)


# --------------------------------------------------------------------- 行内解析
_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def _add_runs(par, text: str) -> None:
    """把 **粗体** 和 `code` 拆成 run，其余按普通文本。"""
    text = text.replace("&nbsp;", " ")
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            r = par.add_run(piece[2:-2])
            r.bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(piece)


def _is_sep_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells if c.strip() != "") \
        and any(c.strip() for c in cells)


def _split_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def convert(md_path: str, docx_path: str) -> int:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    with open(md_path, "r", encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    # 正文中文字体
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    try:
        from docx.oxml.ns import qn
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    except Exception:  # noqa: BLE001
        pass

    i = 0
    n_tables = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ---- 代码块 ----
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue

        # ---- 表格 ----
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            header = _split_row(stripped)
            sep = _split_row(lines[i + 1])
            if _is_sep_row(sep):
                body = []
                j = i + 2
                while j < len(lines) and lines[j].strip().startswith("|"):
                    body.append(_split_row(lines[j]))
                    j += 1
                ncol = len(header)
                t = doc.add_table(rows=1, cols=ncol)
                t.style = "Table Grid"
                for k, h in enumerate(header[:ncol]):
                    cell = t.rows[0].cells[k]
                    cell.text = ""
                    par = cell.paragraphs[0]
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _add_runs(par, h)
                    for run in par.runs:
                        run.bold = True
                        run.font.size = Pt(9)
                for row in body:
                    cells = t.add_row().cells
                    for k in range(ncol):
                        val = row[k] if k < len(row) else ""
                        cells[k].text = ""
                        par = cells[k].paragraphs[0]
                        _add_runs(par, val)
                        for run in par.runs:
                            run.font.size = Pt(9)
                doc.add_paragraph()
                n_tables += 1
                i = j
                continue

        # ---- 标题 ----
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(re.sub(r"\*\*|`", "", m.group(2)), level=level)
            i += 1
            continue

        # ---- 分隔线 ----
        if re.fullmatch(r"-{3,}", stripped):
            doc.add_paragraph("─" * 30).alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # ---- 引用 ----
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Pt(24)
            _add_runs(par, " ".join(buf))
            for run in par.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue

        # ---- 列表 ----
        m = re.match(r"^(\s*)([*\-]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            txt = m.group(3)
            style_name = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            par = doc.add_paragraph(style=style_name)
            if indent:
                par.paragraph_format.left_indent = Pt(18 * (indent + 1))
            _add_runs(par, txt)
            i += 1
            continue

        # ---- 空行 ----
        if not stripped:
            i += 1
            continue

        # ---- 普通段落 ----
        par = doc.add_paragraph()
        _add_runs(par, stripped)
        i += 1

    doc.save(docx_path)
    print(f"[OK] {md_path} -> {docx_path}（{n_tables} 张表）")
    return 0


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        pass
    if len(argv) >= 3:
        src, dst = argv[1], argv[2]
    else:
        src = os.path.join(ROOT, "results", "method_analysis.md")
        dst = os.path.join(ROOT, "results", "method_analysis.docx")
    if not os.path.isabs(src):
        src = os.path.join(ROOT, src)
    if not os.path.isabs(dst):
        dst = os.path.join(ROOT, dst)
    return convert(src, dst)


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
